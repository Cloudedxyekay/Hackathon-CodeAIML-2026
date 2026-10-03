from pathlib import Path


def read_xlsx(path: Path) -> str:
    try:
        from openpyxl import load_workbook
    except ImportError:
        return "XLSX extraction unavailable. Install openpyxl from requirements.txt."

    workbook = load_workbook(path, data_only=True, read_only=True)
    lines = []
    for sheet in workbook.worksheets:
        lines.append(f"[sheet {sheet.title}]")
        for row in sheet.iter_rows(values_only=True):
            values = [str(value) if value is not None else "" for value in row]
            if any(values):
                lines.append(" | ".join(values))
    workbook.close()
    return "\n".join(lines)

